"""Writes the same 'ingredient data' Excel file as the site's 'Export Ingredient Data (Excel)' button, straight from the live Pi
(read-only), so a copy can be kept outside the site. Same columns as the site's Import Ingredient Data (Excel).
    python scripts/export-ingredient-data.py [output.xlsx]      (needs the signed-in cookie jar, COOKIE_JAR env var)"""
import datetime, json, os, subprocess, sys
import openpyxl
from openpyxl.styles import Font

API = os.environ.get("NUTRICOST_API_BASE", "http://192.168.0.50:5001") + "/api/ingredients?includeUnlinked=true"
JAR = os.environ.get("COOKIE_JAR", r"C:\Users\JAMESM~1\AppData\Local\Temp\qa_cookies.txt")
COLS = [("Code", None, "ref"), ("Name", "name", "ref"), ("Category", "cat", "ref"), ("Supplier", "supplier", "ref"),
        ("Energy (kJ)", "kj", "n"), ("Energy (kcal)", "kcal", "n"), ("Fat (g)", "fat", "n"), ("Saturates (g)", "sat", "n"),
        ("Carbohydrate (g)", "carb", "n"), ("Sugars (g)", "sugar", "n"), ("Fibre (g)", "fibre", "n"), ("Protein (g)", "protein", "n"), ("Salt (g)", "salt", "n"),
        ("Allergens", "allergens", "list"), ("Allergen Notes", "allergenNotes", "s"), ("Ingredients", "ingredientsList", "s"),
        ("Pack Size", "packSize", "s"), ("Pack Size Number", "packSizeValue", "n"), ("Pack Size Unit", "packSizeUnit", "s"),
        ("Pack Format", "packFormat", "s"), ("Storage Conditions", "storageConditions", "s"), ("Shelf Life", "shelfLife", "s"),
        ("Unit Weight (g)", "unitWeightG", "n"), ("FVN (Yes/No)", "fvn", "yn")]

raw = subprocess.run(["curl", "-s", "-b", JAR, API], capture_output=True).stdout.decode("utf-8")
ings = json.loads(raw)
assert isinstance(ings, list) and len(ings) > 100, "unexpected API answer (signed in?)"
ings.sort(key=lambda i: (len(str(i.get("code") or "")), str(i.get("code") or "")))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Ingredient Data"
ws.append([c[0] for c in COLS])
for c in ws[1]: c.font = Font(bold=True)
for i in ings:
    row = [i.get("code") or ""]
    for h, f, t in COLS[1:]:
        v = i.get(f)
        if t == "list": row.append(", ".join(v or []))
        elif t == "yn": row.append("Yes" if v else "No")
        elif t == "n": row.append("" if v in (None, "") else float(v) if not float(v).is_integer() else int(v))
        else: row.append("" if v is None else str(v))
    ws.append(row)
for col, (h, f, t) in zip("ABCDEFGHIJKLMNOPQRSTUVWX", COLS):
    ws.column_dimensions[col].width = 40 if h in ("Name", "Ingredients", "Allergen Notes") else 12 if h == "Code" else 16
ws.freeze_panes = "B2"
ab = wb.create_sheet("About")
for r in (["NutriCost ingredient data file"], ["One row per ingredient, matched by Code when imported (Import Ingredient Data (Excel))."],
          ["Code, Name, Category and Supplier are for reference only and are never changed by an import."],
          ["A blank cell leaves the site's value alone. Allergens: comma separated from the 14 allergens. FVN: Yes or No."],
          ["Exported", datetime.datetime.now().isoformat(timespec="seconds")], ["Ingredients in this file", len(ings)]):
    ab.append(r)
out = sys.argv[1] if len(sys.argv) > 1 else "NutriCost ingredient data %s.xlsx" % datetime.date.today().isoformat()
wb.save(out)
print("wrote", out, "-", len(ings), "ingredients")
